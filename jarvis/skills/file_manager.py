import os
from datetime import datetime, timedelta
from typing import Dict, List, Tuple
from jarvis.core.models import ExecutionResult
from jarvis.utils.file_indexer import FileIndexer
from jarvis.utils.memory import Memory

class FileManager:
    """
    Intelligent File Manager.
    - Uses local index (FileIndexer).
    - Parses natural language constraints (provided by Brain).
    - Enforces safety (Search -> Confirm -> Open).
    """
    
    def __init__(self, memory: Memory):
        self.indexer = FileIndexer()
        self.memory = memory
        # Lazy index build: avoid long startup scans.
        # The index will refresh automatically on the first file search.

    def handle(self, intent: Dict) -> ExecutionResult:
        """
        Handle 'files' intent from Brain.
        Intent structure:
        {
            "category": "files",
            "args": {
                "action": "search" | "create" | "delete",
                # Search args
                "type": "pdf", "time_range": "yesterday", "name_contains": "invoice",
                # Create/Delete args
                "name": "notes.txt", "content": "..."
            }
        }
        """
        args = intent.get("args", {})
        action = args.get("action", "search") # Default to search
        
        if action == "create":
            return self.create_file(args)
        elif action == "delete":
            return self.delete_file(args)
        
        # Default: Search
        
        # 1. Parse Constraints
        constraints = self._parse_constraints(args)
        
        # 2. Search
        print(f"[FileManager] Searching with constraints: {constraints}")
        results = self.indexer.search(constraints)
        
        # 3. Handle Results
        if not results:
            return ExecutionResult(False, "I couldn't find any matching items.")
            
        if len(results) == 1:
            item = results[0]
            item_type = "folder" if item["ext"] == "folder" else "file"
            self.memory.set_context("file_candidates", [item])
            return ExecutionResult(
                True, 
                f"I found one {item_type}: '{item['name']}'. Should I open it?",
                data={"count": 1, "candidates": results}
            )
            
        # Multiple results
        self.memory.set_context("file_candidates", results)
        
        # Format list for display/TTS
        choices = []
        for i, r in enumerate(results[:5]): # Limit to 5
            item_type = "[DIR]" if r["ext"] == "folder" else "[FILE]"
            choices.append(f"{i+1}. {item_type} {r['name']} ({r['location']})")
        
        choices_str = "\n".join(choices)
        return ExecutionResult(
            True,
            f"I found multiple items:\n{choices_str}\nWhich one should I open?",
            data={"count": len(results), "candidates": results}
        )

    def open_confirmed(self, selection: int) -> ExecutionResult:
        """
        Execute open action after user confirms.
        selection: 1-based index (e.g. "first one" -> 1)
        """
        candidates = self.memory.get_context("file_candidates")
        if not candidates:
            return ExecutionResult(False, "No file selected.", error="NO_CONTEXT")
            
        idx = selection - 1
        if idx < 0 or idx >= len(candidates):
            return ExecutionResult(False, "Invalid selection.")
            
        target = candidates[idx]
        try:
            print(f"[FileManager] Opening: {target['path']}")
            os.startfile(target["path"])
            return ExecutionResult(True, f"Opening {target['name']}")
        except Exception as e:
            return ExecutionResult(False, f"Failed to open file: {e}")

    def _parse_constraints(self, args: Dict) -> Dict:
        """Convert natural language args to indexer constraints."""
        c = {}
        
        # Type
        if "type" in args:
            c["type"] = args["type"] # ".pdf", "pdf", etc handled by indexer
            
        # Time Parser
        # "yesterday", "today", "last week"
        tr = args.get("time_range")
        if tr:
            now = datetime.now()
            today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
            
            if tr == "today":
                c["time_range"] = {"start": today_start, "end": now}
            elif tr == "yesterday":
                start = today_start - timedelta(days=1)
                end = today_start - timedelta(seconds=1)
                c["time_range"] = {"start": start, "end": end}
            elif tr == "last week":
                start = today_start - timedelta(days=7)
                c["time_range"] = {"start": start, "end": now}
        
        # Location ("downloaded" -> Downloads)
        action = args.get("action", "")
        if "downloaded" in action:
            c["locations"] = ["Downloads"]
        elif "desktop" in action:
             c["locations"] = ["Desktop"]
             
        # Name
        if "name_contains" in args:
            c["name_contains"] = args["name_contains"]
            
        return c

    # --- New Methods from files.py ---

    def create_file(self, details: Dict) -> ExecutionResult:
        """Create a blank/template file."""
        from pathlib import Path
        import subprocess
        
        name = details.get("name", f"file_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt")
        # Handle implied types if name has no extension
        if "." not in name:
            if details.get("type") == "word": name += ".docx"
            elif details.get("type") == "pdf": name += ".pdf"
            elif details.get("type") == "ppt": name += ".pptx"
            else: name += ".txt"
            
        downloads = Path.home() / "Downloads"
        downloads.mkdir(exist_ok=True)
        filepath = downloads / name
        
        try:
            if name.endswith(".docx"):
                return self._create_word(filepath)
            elif name.endswith(".pdf"):
                return self._create_pdf(filepath)
            elif name.endswith(".pptx"):
                return self._create_ppt(filepath)
            else:
                filepath.touch()
                try:
                    subprocess.Popen(["notepad.exe", str(filepath)])
                except: pass
                return ExecutionResult(True, f"Created {name} in Downloads", data={"path": str(filepath)})
        except Exception as e:
            return ExecutionResult(False, f"Failed to create file: {e}")

    def delete_file(self, details: Dict) -> ExecutionResult:
        """Delete a file (Safety: Only Downloads for now)."""
        from pathlib import Path
        
        name = details.get("name")
        if not name:
             return ExecutionResult(False, "Which file should I delete?")
             
        # Safety: Require explicit confirmation in the command or a 2-step process
        # But for now, we'll check if 'confirm' is in details (passed by Brain?)
        if not details.get("confirm"):
             return ExecutionResult(False, f"Please say 'delete {name} confirm' to safely delete it.")
             
        downloads = Path.home() / "Downloads"
        target = downloads / name
        
        if not target.exists():
             # Try fuzzy match?
             return ExecutionResult(False, f"File {name} not found in Downloads.")
             
        try:
            target.unlink()
            return ExecutionResult(True, f"Deleted {name}")
        except Exception as e:
            return ExecutionResult(False, f"Failed to delete: {e}")

    # --- Helpers ---
    def _create_word(self, path) -> ExecutionResult:
        try:
            from docx import Document
            import subprocess
            doc = Document()
            doc.add_heading("Document", 0)
            doc.save(str(path))
            subprocess.Popen(["start", str(path)], shell=True)
            return ExecutionResult(True, f"Created Word doc: {path.name}")
        except ImportError: return ExecutionResult(False, "Install python-docx")

    def _create_pdf(self, path) -> ExecutionResult:
        try:
            from reportlab.pdfgen import canvas
            import subprocess
            c = canvas.Canvas(str(path))
            c.drawString(100, 750, "Created by JARVIS")
            c.save()
            subprocess.Popen(["start", str(path)], shell=True)
            return ExecutionResult(True, f"Created PDF: {path.name}")
        except ImportError: return ExecutionResult(False, "Install reportlab")

    def _create_ppt(self, path) -> ExecutionResult:
        try:
            from pptx import Presentation
            import subprocess
            prs = Presentation()
            prs.save(str(path))
            subprocess.Popen(["start", str(path)], shell=True)
            return ExecutionResult(True, f"Created PPT: {path.name}")
        except ImportError: return ExecutionResult(False, "Install python-pptx")
